"""
Textora AI - PowerPoint PPTX Document Extractor
Extracts slide titles, body text shapes, slide tables, and presenter notes
with preserved slide numbering and section breakdown.
"""
import io
from typing import Dict, Any, List
from pptx import Presentation

def extract_pptx(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    try:
        prs = Presentation(io.BytesIO(file_bytes))
        sections: List[Dict[str, Any]] = []
        full_text_parts: List[str] = []
        
        for idx, slide in enumerate(prs.slides):
            slide_num = idx + 1
            slide_title = ""
            body_texts: List[str] = []
            
            # Extract slide title if available
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()
            
            # Extract shapes, text boxes, tables
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue  # already captured
                    
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            body_texts.append(text)
                            
                elif shape.has_table:
                    for row in shape.table.rows:
                        row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_vals:
                            body_texts.append(" | ".join(row_vals))
                            
            # Presenter notes
            notes_text = ""
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_raw = slide.notes_slide.notes_text_frame.text.strip()
                    if notes_raw:
                        notes_text = notes_raw
            except Exception:
                pass
                
            title_display = slide_title if slide_title else f"Slide {slide_num}"
            
            slide_content_parts = []
            if slide_title:
                slide_content_parts.append(f"Title: {slide_title}")
            if body_texts:
                slide_content_parts.extend(body_texts)
            if notes_text:
                slide_content_parts.append(f"[Presenter Notes]: {notes_text}")
                
            slide_text = "\n".join(slide_content_parts).strip()
            words = len(slide_text.split()) if slide_text else 0
            
            sections.append({
                "section_id": f"slide_{slide_num}",
                "title": f"Slide {slide_num}: {title_display}",
                "slide_number": slide_num,
                "slide_title": title_display,
                "text": slide_text,
                "notes": notes_text,
                "word_count": words,
                "char_count": len(slide_text),
                "is_empty": len(slide_text) == 0
            })
            
            if slide_text:
                full_text_parts.append(f"=== Slide {slide_num}: {title_display} ===\n{slide_text}")
                
        full_text = "\n\n".join(full_text_parts).strip()
        
        return {
            "success": True,
            "filename": filename,
            "file_type": "pptx",
            "slide_count": len(prs.slides),
            "total_words": len(full_text.split()) if full_text else 0,
            "total_chars": len(full_text),
            "sections": sections,
            "full_text": full_text
        }
        
    except Exception as e:
        return {
            "success": False,
            "error": f"Failed to parse PPTX presentation: {str(e)}",
            "sections": [],
            "full_text": "",
            "slide_count": 0
        }
