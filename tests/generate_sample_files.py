"""
Generates sample TXT, DOCX, PPTX, and PDF documents for testing uploads.
"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

TESTS_DIR = Path(__file__).resolve().parent

HUMAN_TEXT = """The Old Harbor Market
Every morning before dawn breaks over the bay, Mr. Callahan arranges his wooden crates of fresh mackerel. He has been running this small fish stall for over forty years, his hands rough and mapped with faint scars from fishing line. 

"The sea doesn't owe you anything," he told me once, tossing a scrap of ice toward a hungry seagull. "You wake up, you cast your net, and you respect whatever comes up."

The harbor smelled of saltwater, creosote, and the sweet caramel from the bakery across the cobbled street. By seven in the morning, delivery trucks were already rumbling down the alleyways, drivers shouting greetings in three different languages. It was chaotic, loud, and thoroughly alive."""

AI_TEXT = """The Transformative Landscape of Modern Education
In today's fast-paced digital world, artificial intelligence plays a crucial role in revolutionizing pedagogical frameworks across educational institutions. Moreover, it is essential to consider the multifaceted nature of automated adaptive learning systems.

This technological evolution serves as a testament to the dynamic landscape of modern computational innovation. Furthermore, machine learning models delve into student performance metrics to unlock personalized instructional pathways.

At its core, this rich tapestry of interconnected digital tools fosters a sense of engagement among diverse learners. In conclusion, it is important to remember that ethical governance and transparent oversight remain paramount as we navigate this transformative pedagogical journey."""

def create_files():
    # 1. TXT
    (TESTS_DIR / "sample_human_essay.txt").write_text(HUMAN_TEXT, encoding="utf-8")
    (TESTS_DIR / "sample_ai_paper.txt").write_text(AI_TEXT, encoding="utf-8")
    
    # 2. DOCX
    doc = Document()
    doc.add_heading("The Transformative Landscape of Modern Education", level=1)
    doc.add_paragraph(AI_TEXT)
    doc.add_heading("Section 2: Methodological Framework", level=2)
    doc.add_paragraph("Furthermore, delving into systemic algorithmic evaluations reveals significant pedagogical enhancements across universities.")
    doc.save(str(TESTS_DIR / "sample_ai_paper.docx"))
    
    # 3. PPTX
    prs = Presentation()
    # Slide 1
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "AI Integration in Academic Research"
    s1.placeholders[1].text = "An Overview of Methodological Innovations"
    if s1.has_notes_slide:
        s1.notes_slide.notes_text_frame.text = "Welcome everyone to today's research briefing."
        
    # Slide 2
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Key Findings and Synthesis"
    s2.placeholders[1].text = (
        "In conclusion, modern natural language processing serves as a testament to algorithmic progress. "
        "Moreover, multifaceted models delve into research archives with unprecedented efficiency."
    )
    if s2.has_notes_slide:
        s2.notes_slide.notes_text_frame.text = "Highlight the statistical significance of slide 2 metrics."
        
    prs.save(str(TESTS_DIR / "sample_presentation.pptx"))
    
    # 4. PDF
    pdf_path = str(TESTS_DIR / "sample_document.pdf")
    pdf_doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("<b>Textora AI Test Document: Academic Writing</b>", styles["Title"]),
        Spacer(1, 15),
        Paragraph(AI_TEXT, styles["Normal"]),
        Spacer(1, 15),
        Paragraph("<b>Page 2 Considerations:</b>", styles["Heading2"]),
        Paragraph("In the realm of educational technology, it is worth noting that continuous human evaluation underpins trustworthy machine adoption.", styles["Normal"])
    ]
    pdf_doc.build(story)
    print("Sample test files generated in tests/ [OK]")

if __name__ == "__main__":
    create_files()
