"""
Textora AI - Backend Validation Tests
Verifies linguistic engine, document extractors, database persistence, and report generation.
"""
import io
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.engine.analyzer import analyze_text
from app.database import init_db, save_audit, list_audits, get_audit, delete_audit
from app.reports.pdf_generator import generate_pdf_report
from app.extractor.text_extractor import extract_document
from docx import Document
from pptx import Presentation

HUMAN_SAMPLE = (
    "I still remember the scent of salt spray and diesel hanging over the pier that chilly October dawn. "
    "Dad fumbled with the rusted padlock on the tackle box, his knuckles red from the wind. "
    "He swore under his breath when the key jammed, but grinned the minute the brass tumblers clicked into place. "
    "'First cast belongs to the early riser,' he said, handing me the weathered fiberglass rod. "
    "We stood shoulder to shoulder for nearly three hours without landing a single bite. "
    "It didn't matter in the slightest. "
    "The rhythm of the tide against the wooden pilings felt like a steady, unspoken conversation."
)

AI_SAMPLE = (
    "In today's fast-paced digital world, artificial intelligence plays a crucial role in revolutionizing modern industries. "
    "Moreover, it is essential to consider the multifaceted nature of automated machine learning systems across society. "
    "This technological paradigm serves as a testament to the dynamic landscape of modern computational innovation. "
    "Furthermore, machine learning algorithms delve into vast datasets to unlock unprecedented analytical potential. "
    "At its core, this rich tapestry of interconnected technologies fosters a sense of progress across diverse sectors. "
    "In conclusion, it is important to remember that ethical considerations remain paramount as we navigate this transformative journey."
)

def run_tests():
    print("--> Initializing SQLite Database...")
    init_db()
    
    print("--> Testing Engine on Human Sample...")
    human_res = analyze_text(HUMAN_SAMPLE)
    assert human_res["success"], "Human analysis failed"
    print(f"Human Sample AI Score: {human_res['scores']['ai_score']}%, Human Score: {human_res['scores']['human_score']}%, Classification: {human_res['classification']}")
    assert human_res['scores']['human_score'] > 50, "Expected human score to predominate"
    
    print("--> Testing Engine on AI Sample...")
    ai_res = analyze_text(AI_SAMPLE)
    assert ai_res["success"], "AI analysis failed"
    print(f"AI Sample AI Score: {ai_res['scores']['ai_score']}%, Human Score: {ai_res['scores']['human_score']}%, Classification: {ai_res['classification']}")
    assert ai_res['scores']['ai_score'] > 60, "Expected AI score to predominate"
    assert ai_res['metrics']['formulaic_patterns']['total_markers'] >= 4, "Expected formulaic markers detected"
    
    print("--> Testing Database Persistence...")
    audit_id = save_audit("AI Sample Doc", "txt", ai_res)
    assert audit_id is not None
    fetched = get_audit(audit_id)
    assert fetched["document_name"] == "AI Sample Doc"
    assert fetched["word_count"] > 0
    all_audits = list_audits()
    assert len(all_audits) > 0
    
    print("--> Testing PDF Report Generation...")
    pdf_bytes = generate_pdf_report(ai_res, "AI Sample Doc")
    assert len(pdf_bytes) > 1000, "PDF report output is too small"
    print(f"Generated PDF successfully ({len(pdf_bytes)} bytes)")
    
    print("--> Testing DOCX Creation & Extraction...")
    doc = Document()
    doc.add_heading("Market Analysis Overview", level=1)
    doc.add_paragraph("This is a synthesized test document designed to evaluate paragraph extraction.")
    doc_io = io.BytesIO()
    doc.save(doc_io)
    docx_res = extract_document(doc_io.getvalue(), "sample.docx")
    assert docx_res["success"]
    assert len(docx_res["sections"]) >= 1
    assert "Market Analysis Overview" in docx_res["full_text"]
    print("DOCX Extraction: PASS")
    
    print("--> Testing PPTX Creation & Extraction...")
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Quarterly Business Review"
    slide.placeholders[1].text = "Revenue grew by 24% year-over-year."
    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    pptx_res = extract_document(pptx_io.getvalue(), "sample.pptx")
    assert pptx_res["success"]
    assert len(pptx_res["sections"]) == 1
    assert "Quarterly Business Review" in pptx_res["full_text"]
    print("PPTX Extraction: PASS")
    
    print("--> Cleaning test record...")
    delete_audit(audit_id)
    
    print("\nALL BACKEND TESTS PASSED SUCCESSFULLY! [OK]")

if __name__ == "__main__":
    run_tests()
